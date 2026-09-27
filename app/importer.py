"""上传文件解析：CSV（UTF-8/GBK 自动识别）与 XLSX；兼容修正模板列。"""
import csv
from dataclasses import dataclass, field
from datetime import datetime, date
from pathlib import Path

# 标准列 -> 可识别的表头别名（小写比较）
COLUMN_ALIASES = {
    "name":        {"姓名", "学员姓名", "name"},
    "id_number":   {"证件号", "证件号码", "身份证号", "id_number", "idnumber"},
    "course_code": {"课程代码", "课程编号", "course_code"},
    "course_name": {"课程名称", "course_name"},
    "hours":       {"学时", "学时数", "hours"},
    "train_date":  {"培训日期", "日期", "train_date"},
}

# 修正模板附加列：解析时忽略，且其出现意味着这是修正模板
EXTRA_TEMPLATE_COLUMNS = {"原始行号", "异常来源", "错误码", "错误原因"}

REQUIRED_FIELDS = ["name", "id_number", "course_code", "course_name", "hours", "train_date"]


class ParseError(Exception):
    pass


@dataclass
class ParseResult:
    rows: list[dict] = field(default_factory=list)
    is_correction: bool = False


def _norm_header(h: str) -> str:
    return (h or "").strip().lstrip("﻿").lower()


def _build_header_map(header: list) -> tuple[dict[str, int], bool]:
    """把表头映射到标准字段；返回 (字段->列号, 是否修正模板)。"""
    alias_lookup = {}
    for std, aliases in COLUMN_ALIASES.items():
        for a in aliases:
            alias_lookup[_norm_header(a)] = std
    extra = {c.lower() for c in EXTRA_TEMPLATE_COLUMNS}
    field_map: dict[str, int] = {}
    is_correction = False
    for idx, raw in enumerate(header):
        norm = _norm_header(_cell_to_str(raw))
        if not norm:
            continue
        if norm in extra:
            is_correction = True
            continue
        std = alias_lookup.get(norm)
        if std and std not in field_map:
            field_map[std] = idx
    return field_map, is_correction


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _rows_from_matrix(matrix: list[list]) -> ParseResult:
    # 定位表头：第一个非空行
    header_idx = next((i for i, r in enumerate(matrix)
                       if any(_cell_to_str(c) for c in r)), None)
    if header_idx is None:
        raise ParseError("文件为空或没有有效表头")
    field_map, is_correction = _build_header_map(matrix[header_idx])
    missing = [f for f in REQUIRED_FIELDS if f not in field_map]
    if missing:
        names = {"name": "姓名", "id_number": "证件号", "course_code": "课程代码",
                 "course_name": "课程名称", "hours": "学时", "train_date": "培训日期"}
        raise ParseError("缺少必要列：" + "、".join(names[f] for f in missing))
    result = ParseResult(is_correction=is_correction)
    # 行号按 Excel 习惯：文件第 1 行为表头，数据从第 2 行起；空行跳过但行号保持准确
    for i, raw_row in enumerate(matrix[header_idx + 1:], start=header_idx + 2):
        if not any(_cell_to_str(c) for c in raw_row):
            continue
        record = {}
        for f in REQUIRED_FIELDS:
            idx = field_map[f]
            record[f] = _cell_to_str(raw_row[idx]) if idx < len(raw_row) else ""
        record["row_no"] = i
        result.rows.append(record)
    return result


def _read_csv(path: Path) -> ParseResult:
    text = None
    for enc in ("utf-8-sig", "gbk"):
        try:
            text = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ParseError("无法识别文件编码（仅支持 UTF-8 / GBK）")
    return _rows_from_matrix(list(csv.reader(text.splitlines())))


def _read_xlsx(path: Path) -> ParseResult:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise ParseError("服务器未安装 openpyxl，无法解析 xlsx") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    matrix = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return _rows_from_matrix(matrix)


def parse_file(path: str | Path) -> ParseResult:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    raise ParseError(f"不支持的文件格式：{suffix}（仅支持 .csv / .xlsx）")
